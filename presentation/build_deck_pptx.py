#!/usr/bin/env python3
"""Build the editable PowerPoint deck for NN-Accel-Triage (academic, 16:9)."""
from __future__ import annotations
import pathlib
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

ROOT = pathlib.Path("/home/user/nn-accel-triage")
OUT = ROOT / "presentation" / "nn_accel_triage.pptx"
GIF = ROOT / "presentation" / "assets" / "pipeline_animation.gif"
IMG_VERIFY = ROOT / "reports" / "dash_verification.png"
IMG_EXPLAIN = ROOT / "reports" / "dash_explain_debug.png"
IMG_PACK = ROOT / "reports" / "dash_packaging.png"

# palette
INK = RGBColor(0x1A, 0x20, 0x2C)
MUTE = RGBColor(0x5A, 0x6B, 0x82)
ACCENT = RGBColor(0x2B, 0x6C, 0xB0)
ACCENT2 = RGBColor(0x2F, 0x85, 0x5A)
LIGHT = RGBColor(0xF7, 0xF9, 0xFC)
PANEL2 = RGBColor(0xEE, 0xF2, 0xF7)
BORDER = RGBColor(0xD6, 0xDE, 0xE8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x14, 0x1B, 0x2D)
WARN = RGBColor(0xB7, 0x79, 0x1F)
BAD = RGBColor(0xC5, 0x30, 0x30)
GREENBG = RGBColor(0xE6, 0xF4, 0xEC)
WARNBG = RGBColor(0xFB, 0xF1, 0xDD)
FONT = "Arial"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def slide(bg=WHITE):
    s = prs.slides.add_slide(BLANK)
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    r.fill.solid(); r.fill.fore_color.rgb = bg; r.line.fill.background()
    r.shadow.inherit = False
    s.shapes._spTree.remove(r._element); s.shapes._spTree.insert(2, r._element)
    return s


def _set_font(run, size, color, bold=False, italic=False, font=FONT):
    run.font.size = Pt(size); run.font.color.rgb = color
    run.font.bold = bold; run.font.italic = italic; run.font.name = font


def text(s, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         space_after=4, line=1.05):
    tb = s.shapes.add_textbox(x, y, w, h); tf = tb.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    if isinstance(runs[0], tuple): runs = [runs]
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(space_after); p.space_before = Pt(0)
        p.line_spacing = line
        for seg in para:
            txt, size, color = seg[0], seg[1], seg[2]
            kw = seg[3] if len(seg) > 3 else {}
            r = p.add_run(); r.text = txt
            _set_font(r, size, color, kw.get("b", False), kw.get("i", False),
                      kw.get("font", FONT))
    return tb


def plain_table(tbl):
    """Strip the default (dark-banded) table style so explicit cell fills show."""
    tbl.first_row = False
    tbl.horz_banding = False
    tblPr = tbl._tbl.tblPr
    for child in list(tblPr):
        if child.tag == qn('a:tableStyleId'):
            tblPr.remove(child)


def rrect(s, x, y, w, h, fill, line=None, lw=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    sp = s.shapes.add_shape(shape, x, y, w, h)
    sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line; sp.line.width = Pt(lw)
    sp.shadow.inherit = False
    try: sp.adjustments[0] = 0.06
    except Exception: pass
    return sp


def chrome(s, kicker, title, footer_r, title_pt=30):
    text(s, Inches(0.62), Inches(0.42), Inches(11), Inches(0.3),
         [[(kicker.upper(), 12.5, ACCENT, {"b": True})]])
    text(s, Inches(0.62), Inches(0.72), Inches(12.1), Inches(1.0),
         [[(title, title_pt, INK, {"b": True})]], line=1.05)
    ln = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.62), Inches(1.52),
                            Inches(12.1), Pt(2.4))
    ln.fill.solid(); ln.fill.fore_color.rgb = ACCENT; ln.line.fill.background()
    ln.shadow.inherit = False
    # footer
    fl = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.62), Inches(7.0),
                            Inches(12.1), Pt(1))
    fl.fill.solid(); fl.fill.fore_color.rgb = BORDER; fl.line.fill.background()
    fl.shadow.inherit = False
    text(s, Inches(0.62), Inches(7.07), Inches(6), Inches(0.3),
         [[("NN-Accel-Triage", 10.5, MUTE)]])
    text(s, Inches(6.72), Inches(7.07), Inches(6), Inches(0.3),
         [[(footer_r, 10.5, MUTE)]], align=PP_ALIGN.RIGHT)


def bullets(s, x, y, w, items, size=17, gap=8):
    tb = s.shapes.add_textbox(x, y, w, Inches(4.5)); tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0
    for i, segs in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap); p.line_spacing = 1.16
        # bullet marker
        r0 = p.add_run(); r0.text = "▪  "
        _set_font(r0, size, ACCENT, True)
        for seg in segs:
            r = p.add_run(); r.text = seg[0]
            kw = seg[2] if len(seg) > 2 else {}
            _set_font(r, size, seg[1], kw.get("b", False), kw.get("i", False))
    return tb


# ---------------------------------------------------------------- 0 TITLE
s = slide(DARK)
# subtle accent bar
bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.22), SH)
bar.fill.solid(); bar.fill.fore_color.rgb = ACCENT; bar.line.fill.background(); bar.shadow.inherit = False
text(s, Inches(0.85), Inches(1.55), Inches(11), Inches(0.4),
     [[("NEURAL-NETWORK ACCELERATOR VERIFICATION", 14, RGBColor(0x8F, 0xB4, 0xE6), {"b": True})]])
text(s, Inches(0.82), Inches(2.05), Inches(11.5), Inches(1.4),
     [[("NN-Accel-Triage", 58, WHITE, {"b": True})]])
text(s, Inches(0.85), Inches(3.75), Inches(10.5), Inches(1.2),
     [[("Agentic AI for Failure Triage and Root-Cause Hinting", 23, RGBColor(0xC9, 0xD6, 0xEA))],
      [("in Neural-Network Accelerator Verification", 23, RGBColor(0xC9, 0xD6, 0xEA))]], line=1.25)
# badges
bx = Inches(0.85)
for label, val in [("faster triage", "342.6×"), ("triage accuracy", "100%"), ("tests passing", "123")]:
    bb = rrect(s, bx, Inches(5.35), Inches(2.75), Inches(0.62),
               RGBColor(0x1E, 0x2A, 0x44), RGBColor(0x3A, 0x4A, 0x6A), 1.0)
    tf = bb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = val + "  "; _set_font(r, 16, RGBColor(0x7F, 0xD0, 0xA3), True)
    r2 = p.add_run(); r2.text = label; _set_font(r2, 14, RGBColor(0xE6, 0xEC, 0xF5))
    bx = Emu(int(bx) + int(Inches(2.95)))
text(s, Inches(0.85), Inches(6.55), Inches(11), Inches(0.4),
     [[("Research project presentation  ·  2026  ·  RTL + CocoTB/Verilator + Claude", 13, RGBColor(0x8F, 0xA3, 0xC4))]])

# ---------------------------------------------------------------- 1 INTRODUCTION
s = slide()
chrome(s, "1 · Introduction",
       "Verification produces failures faster than humans can triage", "Introduction", 26)
bullets(s, Inches(0.62), Inches(1.8), Inches(12.1), [
    [("Neural-network accelerators (systolic ", INK), ("MAC arrays", INK, {"b": True}),
     (" + ", INK), ("quantization units", INK, {"b": True}),
     (") are pervasive; functional correctness across tensor shapes, quantization, and memory behavior is safety-critical.", INK)],
    [("Verification runs ", INK), ("large nightly regressions", INK, {"b": True}),
     (". When tests fail, engineers ", INK), ("triage by hand", INK, {"b": True}),
     (" — reading logs, waveforms, scoreboard diffs, and commit history to find the root cause.", INK)],
    [("Manual triage is ", INK), ("slow (~3 hours per regression here), inconsistent, and does not scale", INK, {"b": True}),
     (" as the design and test suite grow.", INK)],
    [("We present ", INK), ("NN-Accel-Triage", ACCENT, {"b": True}),
     (": a layered testbench plus an ", INK), ("agentic AI", INK, {"b": True}),
     (" that clusters failures, explains likely root causes, and recommends next debug steps.", INK)],
], size=18, gap=13)

# ---------------------------------------------------------------- 2 PROBLEM
s = slide()
chrome(s, "2 · Problem Statement",
       "Automate what an engineer does after a failing regression", "Problem Statement", 26)
bullets(s, Inches(0.62), Inches(1.85), Inches(7.0), [
    [("Input: ", INK, {"b": True}),
     ("a regression database of RTL simulation mismatches — logs, scoreboard diffs, traces, config, and commit metadata.", INK)],
    [("Required output: ", INK, {"b": True}),
     ("automatically (a) group failures by root cause, (b) explain each in natural language, and (c) recommend debug steps — reproducibly and faster than manual.", INK)],
], size=18, gap=16)
card = rrect(s, Inches(7.95), Inches(1.85), Inches(4.75), Inches(4.5), LIGHT, BORDER, 1.0)
text(s, Inches(8.2), Inches(2.05), Inches(4.3), Inches(0.4),
     [[("Why it is hard", 18, ACCENT, {"b": True})]])
bullets(s, Inches(8.2), Inches(2.65), Inches(4.25), [
    [("Failure signals are ", INK), ("heterogeneous", INK, {"b": True}),
     (" — mismatch rate, error magnitude, first-divergence cycle, config.", INK)],
    [("The ", INK), ("number of root causes is unknown", INK, {"b": True}), (" a-priori.", INK)],
    [("Explanations must be ", INK), ("domain-grounded", INK, {"b": True}), (", not generic.", INK)],
    [("Needs a ", INK), ("fair, repeatable benchmark", INK, {"b": True}), (" vs the manual baseline.", INK)],
], size=15, gap=11)

# ---------------------------------------------------------------- 3 OBJECTIVES
s = slide()
chrome(s, "3 · Objectives", "Five core objectives", "Objectives", 30)
objs = [
    ("i", "Verify functional correctness", " across tensor shapes, quantization modes, and memory behaviors.", ACCENT),
    ("ii", "Cluster regression failures automatically", " by symptom / root cause.", ACCENT),
    ("iii", "Generate natural-language debug summaries", " and likely root-cause hints.", ACCENT),
    ("iv", "Benchmark triage time", " against manual methods.", ACCENT),
    ("v", "Package a reproducible", " failure-analysis benchmark others can run.", ACCENT2),
]
oy = 1.9
for num, bold, rest, col in objs:
    badge = rrect(s, Inches(0.7), Inches(oy), Inches(0.72), Inches(0.72), col)
    tf = badge.text_frame; tf.word_wrap = False
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = num; _set_font(r, 20, WHITE, True)
    text(s, Inches(1.65), Inches(oy + 0.06), Inches(10.9), Inches(0.7),
         [[(bold, 19, INK, {"b": True}), (rest, 19, RGBColor(0x33, 0x41, 0x5C))]],
         anchor=MSO_ANCHOR.MIDDLE, line=1.1)
    oy += 0.97

# ---------------------------------------------------------------- 4 METHODOLOGY
s = slide()
chrome(s, "4 · Methodology",
       "A seven-stage pipeline: from RTL fault to root-cause report", "Methodology", 25)
stages = [
    ("STAGE 1", "Design & Fault Injection", "Parameterized RTL; 13 fault variants across 8 categories (frozen ground truth)."),
    ("STAGE 2", "Simulation & Logging", "CocoTB + Verilator vs NumPy golden model; append 1 JSON record per test."),
    ("STAGE 3", "Feature Extraction", "13 features / record (+ optional VCD divergence signals)."),
    ("STAGE 4", "Clustering", "DBSCAN over feature vectors + rule-based symptom labels."),
    ("STAGE 5", "Agentic Triage", "Claude + RAG (TF-IDF) + multi-turn debug loop → cause, confidence, steps."),
    ("STAGE 6", "Benchmark & Eval", "Per-label P/R/F1 & accuracy vs ground truth; time & speedup; no-LLM baseline."),
    ("STAGE 7", "Packaging", "Docker one-command pipeline, HTML dashboard, released dataset."),
]
COLW, COLH = Inches(2.92), Inches(1.72)
GAPX = Inches(0.16)
x0 = Inches(0.62); y0 = Inches(1.8); yr2 = Inches(3.9)
positions = []
for i in range(4):
    positions.append((Emu(int(x0) + i * (int(COLW) + int(GAPX))), y0))
for i in range(3):
    positions.append((Emu(int(x0) + i * (int(COLW) + int(GAPX))), yr2))
tints = [WHITE, WHITE, WHITE, RGBColor(0xF3,0xF8,0xFD), RGBColor(0xF3,0xF8,0xFD), RGBColor(0xF3,0xF8,0xFD), GREENBG]
edges = [BORDER, BORDER, BORDER, RGBColor(0xBC,0xD3,0xEA), RGBColor(0xBC,0xD3,0xEA), RGBColor(0xBC,0xD3,0xEA), RGBColor(0xBF,0xE3,0xCF)]
for i, (num, ttl, desc) in enumerate(stages):
    x, y = positions[i]
    rrect(s, x, y, COLW, COLH, tints[i], edges[i], 1.25)
    text(s, Emu(int(x) + int(Inches(0.18))), Emu(int(y) + int(Inches(0.14))),
         Emu(int(COLW) - int(Inches(0.34))), Inches(1.5),
         [[(num, 11, ACCENT, {"b": True})],
          [(ttl, 15.5, INK, {"b": True})],
          [(desc, 12, MUTE)]], space_after=3, line=1.08)
# 4th box note
x, y = positions[3][0], yr2
rrect(s, Emu(int(x0) + 3 * (int(COLW) + int(GAPX))), yr2, COLW, COLH, WHITE, BORDER, 1.25,
      shape=MSO_SHAPE.ROUNDED_RECTANGLE)
xn = Emu(int(x0) + 3 * (int(COLW) + int(GAPX)))
text(s, Emu(int(xn) + int(Inches(0.18))), Emu(int(yr2) + int(Inches(0.14))),
     Emu(int(COLW) - int(Inches(0.34))), Inches(1.5),
     [[("DATA BUS", 11, ACCENT2, {"b": True})],
      [("regression_db.jsonl", 14, INK, {"b": True, "font": "Consolas"})],
      [("Append-only shared log wired through every stage (1→2→3→4→5→6→7).", 12, MUTE)]],
     space_after=3, line=1.08)

# ---------------------------------------------------------------- 5 ANIMATION
s = slide()
chrome(s, "5 · System in Action",
       "Watch one failing test flow through the pipeline", "System in Action", 26)
try:
    pic = s.shapes.add_picture(str(GIF), Inches(0.9), Inches(1.85), width=Inches(11.5))
except Exception as e:
    print("gif embed failed:", e)
text(s, Inches(0.62), Inches(6.55), Inches(12.1), Inches(0.4),
     [[("Animated GIF — plays automatically in Slide Show mode. ", 12.5, MUTE, {"i": True}),
       ("A fault flows: RTL → simulate → mismatch record → cluster → Claude triage → root-cause report.", 12.5, MUTE, {"i": True})]])

# ---------------------------------------------------------------- 6 VISUALS
s = slide()
chrome(s, "6 · Visuals — Triage Dashboard",
       "One dashboard: verification, explanation, and packaging", "Visuals", 25)
cols = [
    ("Verification", IMG_VERIFY, "Cluster summary + 27 records (15 pass / 12 fail)."),
    ("Explanation + Debugging", IMG_EXPLAIN, "Per-cluster likely cause + debug steps."),
    ("Packaging / Benchmark", IMG_PACK, "Stage timing + headline metrics."),
]
cw = Inches(3.95); cx = Inches(0.62)
for title_, img, cap in cols:
    card = rrect(s, cx, Inches(1.75), cw, Inches(5.05), LIGHT, BORDER, 1.0)
    text(s, Emu(int(cx) + int(Inches(0.2))), Inches(1.9), Emu(int(cw) - int(Inches(0.4))), Inches(0.35),
         [[(title_, 15, ACCENT, {"b": True})]])
    try:
        from PIL import Image
        iw, ih = Image.open(img).size
        maxw = int(cw) - int(Inches(0.4)); maxh = int(Inches(3.9))
        scale = min(maxw / iw, maxh / ih)
        pw, ph = int(iw * scale), int(ih * scale)
        px = int(cx) + (int(cw) - pw) // 2
        s.shapes.add_picture(str(img), Emu(px), Inches(2.35), width=Emu(pw), height=Emu(ph))
    except Exception as e:
        print("img fail", e)
    text(s, Emu(int(cx) + int(Inches(0.2))), Inches(6.35), Emu(int(cw) - int(Inches(0.4))), Inches(0.6),
         [[(cap, 11.5, MUTE)]], line=1.1)
    cx = Emu(int(cx) + int(cw) + int(Inches(0.2)))

# ---------------------------------------------------------------- 7 OUTPUTS: dataset
s = slide()
chrome(s, "7 · Outputs — Fault Dataset",
       "13 fault variants mapped to 8 ground-truth categories", "Outputs · Dataset", 24)
rows = [
    ("fault_acc_overflow", "Accumulator 32→16 bit; sums overflow silently", "accumulator_overflow"),
    ("fault_acc_w24 / w20", "Accumulator narrowed to 24 / 20 bit", "accumulator_overflow"),
    ("fault_wrong_sign", "Weight cast signed→unsigned; no sign-extend", "sign_extension_error"),
    ("fault_b_unsigned", "B_reg loses signed qualifier", "sign_extension_error"),
    ("fault_off_by_one", "Outer loop ends at N-2; last column dropped", "loop_boundary_error"),
    ("fault_loop_over", "Spatial loop runs N+1×; OOB extra product", "loop_boundary_error"),
    ("fault_subtract", "Accumulator subtracts instead of adds", "arithmetic_error"),
    ("fault_reset / quant_reset", "Reset polarity inverted; stale outputs", "reset_polarity_error"),
    ("fault_quant_shift_fixed", "Per-channel shift replaced by fixed >>1", "shift_error"),
    ("fault_quant_no_clamp", "INT8 saturation clamp removed; wraps", "saturation_error"),
    ("fault_quant_wrong_sign_zp", "zero_pt added unsigned instead of signed", "zero_point_error"),
]
tbl = s.shapes.add_table(len(rows) + 1, 3, Inches(0.62), Inches(1.75),
                         Inches(12.1), Inches(4.9)).table
tbl.columns[0].width = Inches(3.2); tbl.columns[1].width = Inches(5.6); tbl.columns[2].width = Inches(3.3)
hdr = ["Variant", "Injected bug", "Ground-truth label"]
for j, h in enumerate(hdr):
    c = tbl.cell(0, j); c.text = ""
    p = c.text_frame.paragraphs[0]; r = p.add_run(); r.text = h
    _set_font(r, 12, RGBColor(0x33,0x41,0x5C), True)
    c.fill.solid(); c.fill.fore_color.rgb = PANEL2
for i, (a, b, cparts) in enumerate(rows, 1):
    for j, val in enumerate((a, b, cparts)):
        c = tbl.cell(i, j); c.text = ""
        p = c.text_frame.paragraphs[0]; r = p.add_run(); r.text = val
        mono = (j == 0)
        _set_font(r, 11.5, INK if j < 2 else WARN, bold=(j == 2),
                  font="Consolas" if mono else FONT)
        c.fill.solid(); c.fill.fore_color.rgb = WHITE if i % 2 else RGBColor(0xFA,0xFC,0xFE)
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
        c.margin_top = Pt(2); c.margin_bottom = Pt(2)
# strip default table style banding header color
plain_table(tbl)

# ---------------------------------------------------------------- 8 OUTPUTS: metrics
s = slide()
chrome(s, "7 · Outputs — Evaluation Metrics",
       "Per-label triage evaluation vs ground truth (27 records)", "Outputs · Evaluation", 24)
mrows = [
    ("no_fault", "1.00", "1.00", "1.00", "15"),
    ("accumulator_overflow", "1.00", "1.00", "1.00", "3"),
    ("loop_boundary_error", "1.00", "1.00", "1.00", "3"),
    ("reset_polarity_error", "1.00", "1.00", "1.00", "3"),
    ("sign_extension_error", "1.00", "1.00", "1.00", "3"),
]
tbl = s.shapes.add_table(len(mrows) + 1, 5, Inches(0.62), Inches(1.9),
                         Inches(7.3), Inches(3.2)).table
for w, cw in zip(tbl.columns, [Inches(3.0), Inches(1.1), Inches(1.1), Inches(1.0), Inches(1.1)]):
    w.width = cw
for j, h in enumerate(["Root-cause label", "Prec.", "Recall", "F1", "Supp."]):
    c = tbl.cell(0, j); c.text = ""
    p = c.text_frame.paragraphs[0]; r = p.add_run(); r.text = h
    _set_font(r, 11.5, RGBColor(0x33,0x41,0x5C), True)
    c.fill.solid(); c.fill.fore_color.rgb = PANEL2
for i, row in enumerate(mrows, 1):
    for j, val in enumerate(row):
        c = tbl.cell(i, j); c.text = ""
        p = c.text_frame.paragraphs[0]; r = p.add_run(); r.text = val
        _set_font(r, 12, INK, font="Consolas" if j == 0 else FONT)
        c.fill.solid(); c.fill.fore_color.rgb = WHITE if i % 2 else RGBColor(0xFA,0xFC,0xFE)
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
plain_table(tbl)
# tiles on right
tiles = [("100%", "Overall accuracy", ACCENT2), ("1.000", "Mean confidence", ACCENT),
         ("27", "Records evaluated", INK), ("1.00", "Macro-F1", INK)]
tx, ty = Inches(8.35), Inches(1.9)
for k, (v, l, col) in enumerate(tiles):
    ox = int(tx) + (k % 2) * int(Inches(2.15))
    oy = int(ty) + (k // 2) * int(Inches(1.35))
    rrect(s, Emu(ox), Emu(oy), Inches(2.0), Inches(1.2), LIGHT, BORDER, 1.0)
    text(s, Emu(ox + int(Inches(0.16))), Emu(oy + int(Inches(0.14))), Inches(1.7), Inches(0.6),
         [[(v, 26, col, {"b": True})]])
    text(s, Emu(ox + int(Inches(0.16))), Emu(oy + int(Inches(0.75))), Inches(1.75), Inches(0.4),
         [[(l, 11.5, MUTE)]])
text(s, Inches(0.62), Inches(5.5), Inches(12.1), Inches(0.9),
     [[("Note: two accumulator-width faults (w24 / w20) are ", 14, RGBColor(0x33,0x41,0x5C)),
       ("mathematically latent", 14, INK, {"b": True}),
       (" for N=4 INT8 (max partial sum 64,516 < 2²⁰) and correctly appear as passing — not misclassifications.", 14, RGBColor(0x33,0x41,0x5C))]],
     line=1.25)

# ---------------------------------------------------------------- 9 RESULTS
s = slide()
chrome(s, "8 · Results & Conclusion",
       "Agentic triage: 100% accurate and 342× faster than manual", "Results & Conclusion", 25)
tiles = [("100%", "Triage accuracy (LLM)", ACCENT2), ("342.6×", "Speedup vs manual", ACCENT),
         ("31.5 s", "vs ~180 min manual", INK), ("123", "Automated tests pass", INK)]
tx = Inches(0.62)
for v, l, col in tiles:
    rrect(s, tx, Inches(1.75), Inches(2.85), Inches(1.15), LIGHT, BORDER, 1.0)
    text(s, Emu(int(tx) + int(Inches(0.2))), Inches(1.88), Inches(2.5), Inches(0.6),
         [[(v, 27, col, {"b": True})]])
    text(s, Emu(int(tx) + int(Inches(0.2))), Inches(2.5), Inches(2.6), Inches(0.35),
         [[(l, 12, MUTE)]])
    tx = Emu(int(tx) + int(Inches(3.02)))
# comparison table
cmp = [
    ("Manual (human)", "baseline", "—", "~180 min", "Reference baseline"),
    ("Rule-based (no LLM)", "72.7%", "0.000", "~0.07 s", "Ablation — clustering only"),
    ("LLM-augmented (ours)", "100%", "1.000", "31.5 s", "Full agentic pipeline"),
]
tbl = s.shapes.add_table(4, 5, Inches(0.62), Inches(3.15), Inches(12.1), Inches(2.1)).table
for w, cw in zip(tbl.columns, [Inches(3.0), Inches(1.9), Inches(1.9), Inches(2.1), Inches(3.2)]):
    w.width = cw
for j, h in enumerate(["Method", "Accuracy", "Mean conf.", "Triage time", "Note"]):
    c = tbl.cell(0, j); c.text = ""
    p = c.text_frame.paragraphs[0]; r = p.add_run(); r.text = h
    _set_font(r, 12, RGBColor(0x33,0x41,0x5C), True)
    c.fill.solid(); c.fill.fore_color.rgb = PANEL2
for i, row in enumerate(cmp, 1):
    for j, val in enumerate(row):
        c = tbl.cell(i, j); c.text = ""
        p = c.text_frame.paragraphs[0]; r = p.add_run(); r.text = val
        col = INK; bold = (j == 0 and i == 3)
        if j == 1 and i == 3: col = ACCENT2; bold = True
        if j == 1 and i == 2: col = WARN; bold = True
        _set_font(r, 12.5, col, bold)
        c.fill.solid()
        c.fill.fore_color.rgb = RGBColor(0xF1,0xFA,0xF4) if i == 3 else (WHITE if i % 2 else RGBColor(0xFA,0xFC,0xFE))
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
plain_table(tbl)
text(s, Inches(0.62), Inches(5.5), Inches(12.1), Inches(1.2),
     [[("Conclusion: ", 15, INK, {"b": True}),
       ("the agentic component lifts accuracy 72.7% → 100% and confidence 0 → 1.0, while a fully reproducible Docker pipeline and released dataset let others re-run the benchmark. ", 15, RGBColor(0x33,0x41,0x5C)),
       ("Limitations: ", 15, INK, {"b": True}),
       ("small 13-variant fault library and a single accelerator design — future work scales both and adds waveform-grounded retrieval.", 15, RGBColor(0x33,0x41,0x5C))]],
     line=1.3)

# ---------------------------------------------------------------- 10 REFERENCES
s = slide()
chrome(s, "9 · References", "References", "References", 30)
refs_l = [
    "W. Snyder et al. Verilator — open-source SystemVerilog simulator. veripool.org.",
    "cocotb — Coroutine-based Cosimulation Testbench. cocotb.org.",
    "M. Ester, H.-P. Kriegel, J. Sander, X. Xu. “A Density-Based Algorithm for Discovering Clusters (DBSCAN).” KDD, 1996.",
    "F. Pedregosa et al. “scikit-learn: Machine Learning in Python.” JMLR 12, 2011.",
    "P. Lewis et al. “Retrieval-Augmented Generation for Knowledge-Intensive NLP Tasks.” NeurIPS, 2020.",
]
refs_r = [
    "G. Salton, C. Buckley. “Term-Weighting Approaches in Automatic Text Retrieval.” IP&M, 1988.",
    "Anthropic. “Claude Models & Messages API.” Technical documentation, 2024–2026.",
    "N. Jouppi et al. “In-Datacenter Performance Analysis of a Tensor Processing Unit.” ISCA, 2017.",
    "Accellera. Universal Verification Methodology (UVM) 1.2 Standard, 2014.",
    "B. Jacob et al. “Quantization and Training of NNs for Efficient Integer-Arithmetic Inference.” CVPR, 2018.",
]
def reflist(x, items, start):
    tb = s.shapes.add_textbox(x, Inches(1.85), Inches(5.9), Inches(4.6)); tf = tb.text_frame
    tf.word_wrap = True
    for i, t in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(12); p.line_spacing = 1.12
        r = p.add_run(); r.text = f"[{start+i}]  "; _set_font(r, 13, ACCENT, True)
        r2 = p.add_run(); r2.text = t; _set_font(r2, 13, RGBColor(0x33,0x41,0x5C))
reflist(Inches(0.62), refs_l, 1)
reflist(Inches(6.85), refs_r, 6)
text(s, Inches(0.62), Inches(6.6), Inches(12), Inches(0.35),
     [[("Related-work references compiled for context — verify / replace specifics before final submission.", 11.5, MUTE, {"i": True})]])

prs.save(str(OUT))
print("saved", OUT, OUT.stat().st_size // 1024, "KB", "| slides:", len(prs.slides._sldIdLst))
